from pathlib import Path
import json
import shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path('/home/ubuntu/world_character_library')
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
BASE_JSON = OUT / 'world_character_library_data.json'

with BASE_JSON.open('r', encoding='utf-8') as f:
    payload = json.load(f)

worlds = payload['original_worlds']
characters = payload['original_characters']
ips = payload['ip_references']
visuals = payload['visual_assets']
licenses = payload['license_matrix']
relations = payload['relationship_patterns']

world_updates = {
    'W-07': ['W-07','鲸落便利店','完全原创','温馨生活；都市奇幻；幽灵经济','午夜便利店随鲸骨街漂移，顾客用未说出口的话支付；一笔沉默正试图买走整座城市','沉默存量 54；街区稳定度 63；失联顾客 7；互助基金 12,400；陌生人指数 44','归潮便利店；鲸骨街；失物冷柜；潮退地铁站；凌晨裁缝铺；热线旧机房','归潮夜班组；幽灵配送站；街区租赁会；市政静谧计划；白昼商会','疗愈 vs 表达；商业续租 vs 社区归属；安静 vs 求救','无声促销夜；欠话收据；七件失物','高','VIS-O-07'],
    'W-08': ['W-08','祈雨列车','完全原创','民俗奇幻；旅行；气候寓言','一列以乘客放下的执念换雨的列车横穿干旱大陆；它开始拒绝收票','云压 41；全线储水 29；雨票 318；未寄信件 73；神契约争议 4','春雷号；干涸总站；神社货场；蘑菇田避难站；旧水库剧场；云顶维修库','列车局；地方神联席会；气象商会；站台自救会；留影团','个体记忆 vs 公共气候；地方传统 vs 流动权','反向车票；第九节车厢；无轨降雨','高','VIS-O-08'],
    'W-09': ['W-09','第七次小行星婚礼','完全原创','浪漫喜剧；外交科幻；时间分岔','政治婚礼已爆炸六次；第七轮中每个承诺都向公共资源与身份扩散','第 7 次重办；庆典可信度 38；分岔残留 26；媒体热度 89；供氧 67','微重力花园；外交通道；礼服修复舱；媒体环廊；旧矿井教堂；观星殡仪馆','婚礼委员会；外交团；反殖民联盟；小行星矿工会；七秒后平台','明确同意 vs 政治婚约；真实自我 vs 时间线残留','七份誓言；第八轮爆炸；未出现的名字','高','VIS-O-09'],
    'W-10': ['W-10','三重城墙','完全原创','战争；城市建设；阶层政治','城市每十年向外修一圈墙，墙外居民自动失去公民权；本轮石料不足','石料配额 58%；内城信任 46；外墙人口 34,000；边境威胁 51；十年钟 281 天','内城议会厅；旧墙市场；新墙工地；墙外诊所；石灰井；第四门','建设署；城防军；墙外联合会；档案局；石匠工会；第四门网络','安全扩张 vs 公民权；基建效率 vs 归属','缺口地图；古澡堂；多孔城墙','高','VIS-O-10'],
}

for i, world in enumerate(worlds):
    if world[0] in world_updates:
        worlds[i] = world_updates[world[0]]

worlds.extend([
    ['W-11','倒悬图书馆','完全原创','学术奇幻；记忆灾害；城邦治理','倒挂峡谷的图书馆保存尚未发生的故事；书中结局正坠入现实','预叙泄露率 18；悬索磨损 39；禁书缺页 11；来访者 2,800；索引公信力 61','倒悬主馆；峡谷市集；风梯宿舍；封存儿童区；断页渡口；未写城','馆务院；抄写员公社；预言保险行；断页盗团；未写城委员会','预知 vs 自主；知识开放 vs 信息伤害','批注羽笔；坠落绘本；无书周','高','VIS-O-11'],
    ['W-12','玻璃雨林公约','完全原创','生物朋克；环境法庭；跨国阴谋','法定人格雨林要求迁徙；人类必须决定谁有权替非人生命说话','意志清晰度 33；穹顶裂隙 9；支持国 14；孢子半径 23km；就业依赖 71','玻璃穹顶主林；根系议会厅；漂浮宿舍城；种子法庭；河流边检站；雨季学校','雨林信托；矿物联合体；地方护林联盟；生物法庭；孢子信徒；难民合作社','非人权利 vs 人类代理；修复 vs 社区驱逐','根系听诊器；植物引渡；48 小时庭审','高','VIS-O-12'],
    ['W-13','候鸟法庭','完全原创','历史幻想；法律悬疑；多世代旅程','候鸟跨境时审理被史书抹去的旧案；判决要求归还从未出生的王后','史实稳定度 57；迁徙窗口 23天；待审旧案 43；戒严 2国；证人遗忘率 31','候鸟法庭；旧帝国邮驿；边检城；盐湖墓园；空白年鉴库；庭审列车','候鸟书记官；边境法务军；无谱系家庭联盟；复辟史家；墨羽人','历史修复 vs 现实伤害；血缘国界 vs 公民身份','回巢印章；少年逃兵；无判决季','高','VIS-O-13'],
    ['W-14','月面凌晨四点','完全原创','近未来科幻；劳工悬疑；月面社区','每人每日仅有 47 分钟私人氧气；事故录像在所有头盔里循环播放','氧气自治率 18；矿尘暴 3天后；夜班出勤 84；调查可信度 21；地球股价 76','晨昏穹顶；四点食堂；白噪矿井；轨道候车环；私氧菜园；信号墓园','月面开采公司；夜班工会；菜园合作社；地球股东团；儿童互助网','生存配给 vs 劳动权；公司叙事 vs 事故真相','47 分钟密钥；全城停工；系统罢工','高','VIS-O-14'],
    ['W-15','深海中学的失踪夏令营','完全原创','校园悬疑；海洋科幻；成长群像','深沟夏令营归来 23 人；每个孩子都记得自己在两支不同队伍中','回归学生 23；记忆版本 2；水压安全 72；家长信任 39；深沟信号 17','海底穹顶校舍；潮汐体育馆；鲸语实验室；沉船图书馆；海沟营地；上浮检疫舱','学校董事会；学生潜水社；海洋科研队；家属会；海沟观察者；地面媒体','学生叙述权 vs 成人风险控制；科学发现 vs 未成年人保护','蓝压徽章；24 人合影；海沟交换','高','VIS-O-15'],
    ['W-16','十万盏灯的河','完全原创','大河文明；神祇政治；多城邦协作','洪峰夜十万盏灯逆流；九城收到来自未来后代的求救信','河运流量 63；灯愿偏移 22；九城联盟 47；洪峰 19天；信件可信度 11','九城灯港；浮桥菜市；河神档案庙；淤泥剧场；旧堤军营；逆流峡','九城联盟；河工行会；灯火寺；河上商盟；逆流者；信件保管局','祖先/后代 vs 现世；流域公共资源 vs 城邦竞争','逆灯；第十城税法；河流议会','高','VIS-O-16']
])

characters.extend([
    ['C-26','W-07','罗知夏','归潮夜班店长','买下便利店经营权','失踪弟弟把整段求救话语卖给了便利店','库存、夜班排班、顾客信任','不向弱者出售沉默','谢绵：依赖/不信任；余弦：热线同盟','租约剩余 30 天'],
    ['C-27','W-07','谢绵','幽灵配送员','完成死前最后一份包裹','曾配送静谧计划试用沉默，牵连知夏弟弟','夜间路线、幽灵网络','不再替机构运送人类情绪','知夏：赎罪债；傅未言：假租约','第七次送件'],
    ['C-28','W-07','傅未言','街区租赁会代表','把鲸骨街纳入合法经营','能听懂鲸骨，知道它想留下被城市抛弃的人','合同、物业、白昼商会','不让父亲拆潮退站','余弦：旧情；谢绵：债主','无声促销夜提案'],
    ['C-29','W-07','余弦','夜间热线志愿者','重启全城心理热线','曾卖掉悲伤样本，正忘记爱过谁','旧机房、录音库、热线名单','不伪造求助者同意','傅未言：旧情；知夏：信息互赖','03:33 不存在来电'],
    ['C-30','W-08','梁秋','春雷号副车长','让列车准点抵达无雨市','从未放下执念，靠伪票维持自身','行车权限、乘客名册','不拒载病人','何萝：相互监督；万羡：救命债','云压低于 20'],
    ['C-31','W-08','何萝','地方神使者','收回被列车夺走的雨权','家族伪造过神谕，造成一站被弃','仪式、通行文书、神社网络','不把地方神当观光资源','万羡：旧怨；杜莲：祭典冲突','第九车厢交换名字'],
    ['C-32','W-08','万羡','气象商会工程师','将降雨变成稳定服务','发现雨会抹去战争记忆，压着报告不发','模型、资金、维修权限','不伤害旅行中的孩子','梁秋：旧恩；何萝：学术冲突','无雨市前 18 小时报告'],
    ['C-33','W-08','杜莲','站台自救会领队','拿回祖传水渠控制权','水渠早已无水，争夺是维持社区筹码','站台民众、仓库、地方记忆','不以水换取弱者通行权','何萝：旧怨；父亲：记忆被扣留','继承人谈判'],
    ['C-34','W-09','至真','小行星继承人、新娘','完成不牺牲任何人的和平协议','她参与制造第八轮爆炸以保留选择','继承权、礼仪席、公众好感','不在缺乏同意下结婚','罗莎：时间线残留情感；索郁：人格承认','直播前 9 小时'],
    ['C-35','W-09','阿德','婚礼策划师','让本轮不再爆炸','第二轮为救一人牺牲过矿工','流程表、宴会控制、媒体关系','不再拿性命优化庆典','索郁：知道其记忆；至真：雇佣/共谋','改座位表会增残留'],
    ['C-36','W-09','罗莎','矿工会谈判人','为矿工拿到签字权','未来版本曾与至真相爱，不知是否真心','罢工、矿井证据、工人信任','不以工人生命换席位','至真：未说出口的亲密；索郁：债务','氧气低于 45'],
    ['C-37','W-09','索郁','外交团代表','让婚约绑定资源条约','是被删除原订婚对象的复制版本','外交担保、访问权限、安保','不接受被当替身','阿德：愧疚；至真：人格承认','必须选择外交立场'],
    ['C-38','W-10','陆阙','建设署总建筑师','按时完成新墙','把隐蔽水管画进正式图纸','工程图、石匠调度、技术信任','不把住房当政治惩罚','赤城：策略同盟；曹宁：档案风险','公开图纸前夕'],
    ['C-39','W-10','曹宁','墙外诊所医生','维持无证居民医疗通道','内城出生，主动放弃身份寻找失散妹妹','诊所、药网、病患信任','不按身份排序救治','晏饮：档案互保；赤城：旧识','水管破裂前 6 天'],
    ['C-40','W-10','赤城','城防军指挥','加速修墙并避免屠杀','暗中制造袭击以逼议会谈判','军队、边境情报、戒严权','不对平民开火','陆阙：不知情盟友；曹宁：共同童年','袭击证据泄露'],
    ['C-41','W-10','晏饮','档案局学徒','成为正式测绘师','每十年会在户籍中消失，系第四门后代','缺口地图、档案钥匙、学徒网','不让名字被抹去','曹宁：保护者；陆阙：图纸依赖','十年钟归零'],
    ['C-42','W-11','顾棠','高级索引师','维持预叙书目准确','曾删去自己的婚姻预叙，正以绘本回归','索引权、馆务院声望','不替人删除死亡预告','黎铸：师徒互瞒；岩声：共同危机','悬索磨损到 60'],
    ['C-43','W-11','黎铸','风梯快递员、断页线人','让未写城取得开放权','本体是一本未写完的书','风梯路线、盗团、平民网','不让儿童书成为预测商品','顾棠：可写其结局；乔褐：投保关系','被归档即失去行动'],
    ['C-44','W-11','乔褐','预言保险行精算师','收购禁书缺页','患被未来看不见的疾病','金融合同、查询权限、访客名单','不卖亲属预言','黎铸：客户；妹妹：不阅读联盟','收购截止'],
    ['C-45','W-11','岩声','封存儿童区照护员','找回每日消失的员工','是逐渐失名员工的集合体','儿童区钥匙、绘本感应','不让孩子承担预言恐惧','黎铸：能看真实页码；顾棠：共同隐瞒','绘本第 11 页'],
    ['C-46','W-12','姜弥','生物法庭律师','为雨林争取独立代理资格','母亲主导过神经记录计划','程序、证据、国际媒体','不替无语言者伪造意愿','莱娜：互需/争论；周疏：翻译质疑','48 小时庭审'],
    ['C-47','W-12','莱娜·维溪','护林联盟导航员','阻止穹顶扩张','家族以医疗换取记录计划，雨林先想迁走她的村落','林间路线、社区信任、野外技能','不卖祖地','本森：旧识/对立；姜弥：法律互赖','孢子越境'],
    ['C-48','W-12','周疏','孢子信徒翻译者','让人类听懂雨林','听诊器令他失去否定词，翻译无法说不','信徒、仪式、翻译能力','不故意篡改讯号','姜弥：怀疑；雨林：梦中投票','暂停翻译投票'],
    ['C-49','W-12','本森·费因','矿物联合体安保主管','安全封堵穹顶裂隙','资助难民合作社但不能承认，修复会切根系','设备、安保、运输','不朝护林员开枪','莱娜：相识；姜弥：关键证人','12 小时清场令'],
    ['C-50','W-13','闻劭','候鸟书记官','按程序审理王后案','王后是前任书记官为难民创造的法律人格','法条、庭审记录、跨时邮袋','不伪造证人','薛黎：无谱系亲缘；童兵：案卷核心','迁徙窗口 23 天'],
    ['C-51','W-13','塞米·阿尔','边境法务军官','封锁跨时证词列车','祖母被历史删除，担心其回归会让自己消失','戒严权、边检、军法','不遣返儿童证人','薛黎：旧情；童兵：互相识破','戒严延长 5 天'],
    ['C-52','W-13','薛黎','无谱系联盟发言人','让无家谱者拥有完整公民权','由三段女性证词拼合，单一判决可能令其消失','民间档案、演说、家庭网','不让身份证明决定尊严','闻劭：亲缘风险；塞米：旧情','公开听证'],
    ['C-53','W-13','童兵','以少年身体归来的逃兵','找到母亲王后','是王后法律人格受益人，归还王后会令他变未登记死者','军史、身体记忆、古战场路线','不再为君主杀人','塞米：可能后代；闻劭：法理冲突','候鸟转向后老化'],
    ['C-54','W-14','唐回','夜班工会组织者','推动全城 11 分钟停工','事故中为救多数关闭救援舱，死者是伴侣','工会、矿井钥匙、工人声望','不把人留在无氧区','莫岛：地球旧友；雅利：氧源同盟','尘暴前 72 小时'],
    ['C-55','W-14','雅利','私氧菜园培育师','让菜园成为自治氧源','植物疑似替系统选择谁可呼吸','植物、食堂、儿童网','不把氧气用于情感勒索','青笺：秘密监护；唐回：合作','扩张违反安全法'],
    ['C-56','W-14','莫岛','地球股东公关代表','压下事故录像','月面出生后被送地球，身体无法适应重力','媒体、股东、地月通讯','不捏造死者家属采访','唐回：旧友/隐瞒；公司：雇佣压力','股价跌破 50'],
    ['C-57','W-14','青笺','低重力儿童互助网发起人','给每个孩子一块私氧菜地','能与事故系统文字对话，系统称她第二班长','儿童网络、维护孔道、菜园','不泄露同伴家庭氧气记录','雅利：监护；系统：未知同盟','系统请求停矿'],
    ['C-58','W-15','禹蓝','潜水社副社长','回海沟找失踪同学','在珊瑚网留有更勇敢的自己','潜水、学生信任、蓝压徽章','不让同学当科研样本','嘉里：竞争到互保；珊瑚版本：身份冲突','深潜窗 36 小时'],
    ['C-59','W-15','嘉里·帕斯','学生会主席','维持正常开学','为保家人删去学校违规录像','学生会、上浮权限、家长沟通','不公开同学私密记忆','禹蓝：知情者；父亲：董事会压力','听证会 2 天'],
    ['C-60','W-15','许鲸','鲸语实验室研究员','证明信号有语言能力','失踪学生是侄女，早知珊瑚网会放大情绪','实验室、数据、设备','不强制提取孩子记忆','葛衣：旧友决裂；禹蓝：学生信任','24 小时样本令'],
    ['C-61','W-15','葛衣','失踪者家属会发言人','带所有孩子上浮','儿子不愿回来，害怕听到其真实记忆','家属组织、媒体、法律顾问','不以孩子照片换流量','许鲸：互责；嘉里：家长压力','私人探视协议'],
    ['C-62','W-16','谢沙','第七城灯匠','让罢工被看作劳工行动','灯芯来自被抹去的家谱纸','灯匠、工坊、引路技术','不让亡者成为燃料','慈夷：少年旧友；东荇：怀疑','洪峰前复工决定'],
    ['C-63','W-16','慈夷','上游城邦外交使者','推进新堤坝','未来信称她会淹没三城','外交、粮仓、水闸','不以救灾粮换婚盟','谢沙：旧友；东荇：政治互用','弟弟在逆流者'],
    ['C-64','W-16','东荇','河工行会调度员','让工人进入九城代表席','可调灯愿偏移，曾救一城害一城断粮','船工、堤坝、物流','不把工人当洪水数字','慈夷：政治互用；谢沙：质疑','灯芯来源曝光'],
    ['C-65','W-16','遥二','未来信件 archivist','证明信件可公共解读','信件来自另一个未来的自己','信件库、统计、河上家庭网','不删除不舒服的未来','慈夷：托付信件；灯火寺：敌对','公开作者身份']
])

ips.extend([
    ['IP-13','Cyberpunk 2077 / Edgerunners','游戏/动画','高密度都市；职业网络','A','职业、债务、身份与街区声望互相影响','临时团队/中间人；亲密关系/生存选择','城市街区；委托链；企业/民间势力','https://www.cyberpunk.net/us/en/','REF-ONLY'],
    ['IP-14','质量效应','游戏','太空团队；外交政治','A','移动基地；个人任务；跨文化共同决策','船员/公共使命；个人忠诚/外交职责','移动基地；航线；多方危机','https://www.ea.com/games/mass-effect/mass-effect-legendary-edition','REF-ONLY'],
    ['IP-15','英雄联盟：双城之战','动画/游戏宇宙','阶层城市；亲属断裂','A','技术、家庭、阶层空间互相施压','手足关系/公共暴力；治理者/创新者','双城空间；技术项目；民间组织','https://www.netflix.com/tudum/articles/arcane-every-league-of-legends-character','REF-ONLY'],
    ['IP-16','最后生还者','游戏/电视剧','末日共同体；照护关系','A','资源压力下的照护、共同体规则与信任','照护者/被照护者；社区领袖/离散亲属','避难点；出行；资源与信息','https://www.hbomax.com/show/93ba22b1-833e-47ba-ae94-8ee7b9eefa9a/cast-and-crew','REF-ONLY'],
    ['IP-17','集合啦！动物森友会','游戏','温馨生活；社区日程','A','真实时间、季节、居民与环境留痕','邻居/组织者；居民/公共空间','岛屿；节日；日程；环境照料','https://animalcrossing.nintendo.com/new-horizons/explore/','REF-ONLY'],
    ['IP-18','上古卷轴 Online','MMO','多地区；职业组织；持续世界','A','区域文化、公会职业、公共事件和路线图','地方居民/公会；新玩家/旧事件','区域入口；季度事件；职业节点','https://www.elderscrollsonline.com/en-us/home','REF-ONLY'],
    ['IP-19','海贼王','漫画/动画/真人剧集','航海群像；地点考验','A','团队成员的个人梦想在不同地点被分别考验','团队成员/个人梦想；共同体/外部挑战','移动团队；目的地；个人任务','https://www.netflix.com/tudum/articles/one-piece-season-2-trailer','REF-ONLY'],
    ['IP-20','饥饿游戏','小说/电影','反乌托邦；公开仪式；阶层','A','公共评价、地区资源和媒体叙事互相强化','当事人/媒体；地区社群/统治结构','公共仪式；区域资源；反叙事网络','https://www.lionsgate.com/franchises/the-hunger-games','REF-ONLY']
])

visuals.extend([
    ['VIS-O-07','原创','鲸落便利店关键视觉','W-07','/visuals/original_w07_whalebone_convenience_keyart.png','16:9；2560×1440','鲸骨街；暖店光；夜班；潮雾','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-08','原创','祈雨列车关键视觉','W-08','/visuals/original_w08_rain_train_keyart.png','16:9；2560×1440','雨云列车；干旱；票根；站台','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-09','原创','小行星婚礼关键视觉','W-09','/visuals/original_w09_asteroid_wedding_keyart.png','16:9；2560×1440','零重力仪式；裂痕请柬；星光','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-10','原创','三重城墙关键视觉','W-10','/visuals/original_w10_three_walls_keyart.png','16:9；2560×1440','同心城墙；工地；公民空间','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-11','原创','倒悬图书馆关键视觉','W-11','/visuals/original_w11_inverted_library_keyart.png','16:9；2560×1440','峡谷；悬索；倒挂书城','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-12','原创','玻璃雨林关键视觉','W-12','/visuals/original_w12_glass_rainforest_keyart.png','16:9；2560×1440','透明穹顶；根系法庭；雨林','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-13','原创','候鸟法庭关键视觉','W-13','/visuals/original_w13_migratory_court_keyart.png','16:9；2560×1440','候鸟；庭审列车；盐湖','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-14','原创','月面凌晨四点关键视觉','W-14','/visuals/original_w14_lunar_4am_keyart.png','16:9；2560×1440','夜班食堂；私氧菜园；尘暴','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-15','原创','深海中学关键视觉','W-15','/visuals/original_w15_deepsea_school_keyart.png','16:9；2560×1440','深海学校；海沟；学生群像','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-16','原创','十万盏灯的河关键视觉','W-16','/visuals/original_w16_river_lanterns_keyart.png','16:9；2560×1440','逆流灯火；九城；洪峰','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['REF-02','第三方','官方视觉/视频研究入口','IP-13 至 IP-20','见 IP_参考库工作表来源链接','网页/视频','仅用于抽象关系、场景、节奏研究','REF-ONLY','不可下载、嵌入、训练或商用；链接不是授权']
])

relations.extend([
    ['RLP-06','共同体财产中的私人债','店铺/车票/氧气/灯火维护者','私人创伤被写入公共价值','贡献 → 控制 → 公开账本/退出','照护、失能与隐性劳动不可被简化为低贡献'],
    ['RLP-07','同一人的多版本关系','时间分岔/共享记忆/档案改写','多个版本主张真实承诺','否认 → 利用 → 多方同意/身份冲突','每个版本均需独立意愿与伤害边界'],
    ['RLP-08','非人行动者的代理困境','雨林/列车/河流/系统/书本','表达渠道可能失真且被代理垄断','代言 → 质疑 → 共治/拒绝翻译','不可让人类代理者垄断非人意志'],
    ['RLP-09','跨墙或跨代的未证亲缘','城市隔离/历史删除/迁徙家庭','法律与档案使亲缘不可证','寻证 → 互保 → 承认/失去旧身份','血缘不是唯一合法家庭形式'],
    ['RLP-10','公共角色的情感回声','外交官/教师/律师/领队','私人失去影响公共决策','克制 → 泄露 → 透明问责/退出','责任是披露与承担后果，不是压抑情感'],
    ['RLP-11','青少年自治与成人风险管理','学生/未成年人/机构人员','关键信息与安全责任分属不同主体','保护 → 排除 → 共同制定边界','明确安全、支持、隐私与非浪漫化原则']
])

change_log = [
    ['v0.2','2026-08-25','世界','孵化 W-07 至 W-10；新增 W-11 至 W-16','10 个高优先级原创世界，均有世界状态、地点、阵营、冲突、秘密、物品与任务钩子。'],
    ['v0.2','2026-08-25','角色','新增 C-26 至 C-65','40 名原创角色，均含目标、秘密、资源、底线、关键关系与个人时钟。'],
    ['v0.2','2026-08-25','关系','新增 RLP-06 至 RLP-11','新增共同体债务、版本人格、非人代理、未证亲缘、情感问责、青少年自治范式。'],
    ['v0.2','2026-08-25','IP 研究','新增 IP-13 至 IP-20','8 个第三方 IP 仅限结构参考；全部带官方来源 URL 与 REF-ONLY 标签。'],
    ['v0.2','2026-08-25','视觉','新增 VIS-O-07 至 VIS-O-16','10 张原创生成关键视觉；第三方官方视听入口仅保留为链接。']
]

payload = {
    'metadata': {
        'library_name': 'World / Character 内容与视觉素材库',
        'version': 'v0.2',
        'updated_at': '2026-08-25 GMT+8',
        'counts': {'original_worlds': len(worlds), 'original_characters': len(characters), 'ip_references': len(ips), 'visual_assets': len(visuals), 'relationship_patterns': len(relations)},
        'license_notice': '第三方 IP 及其官方媒体均为 REF-ONLY；链接不构成许可。'
    },
    'original_worlds': worlds,
    'original_characters': characters,
    'ip_references': ips,
    'visual_assets': visuals,
    'license_matrix': licenses,
    'relationship_patterns': relations,
    'change_log': change_log
}

(OUT / 'world_character_library_data_v0.2.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
(OUT / 'world_character_library_data.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')

wb = Workbook()
wb.remove(wb.active)
header_fill = PatternFill('solid', fgColor='12324A')
subheader_fill = PatternFill('solid', fgColor='1F567D')
accent_fill = PatternFill('solid', fgColor='E7F1F7')
warn_fill = PatternFill('solid', fgColor='FBE9E7')
new_fill = PatternFill('solid', fgColor='E8F5E9')
thin = Side(style='thin', color='C6D4DF')

def add_sheet(name, headers, rows, widths, table_name, link_cols=None, new_id_prefixes=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    for c, value in enumerate(headers, 1):
        cell = ws.cell(1, c, value)
        cell.fill = header_fill
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=thin)
    for r, row in enumerate(rows, 2):
        is_new = bool(new_id_prefixes and str(row[0]) in new_id_prefixes)
        for c, value in enumerate(row, 1):
            cell = ws.cell(r, c, value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=thin)
            cell.fill = new_fill if is_new else (accent_fill if r % 2 == 0 else PatternFill(fill_type=None))
            if link_cols and c in link_cols and isinstance(value, str) and value.startswith('http'):
                cell.hyperlink = value
                cell.style = 'Hyperlink'
        ws.row_dimensions[r].height = 46
    ws.row_dimensions[1].height = 34
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(c)].width = w
    if rows:
        ref = f'A1:{get_column_letter(len(headers))}{len(rows)+1}'
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False)
        ws.add_table(tab)
    return ws

index = wb.create_sheet('00_索引')
index.sheet_view.showGridLines = False
index.merge_cells('A1:D1')
index['A1'] = 'World / Character 内容与视觉素材库 v0.2'
index['A1'].font = Font(size=20, bold=True, color='FFFFFF')
index['A1'].fill = header_fill
index['A1'].alignment = Alignment(horizontal='center', vertical='center')
index.row_dimensions[1].height = 40
summary = [
    ['版本','v0.2 增量扩充；保留 v0.1 内容，不重做旧资产。'],
    ['原创世界','16 个：原有 6 个旗舰世界 + 4 个被完整孵化的种子 + 6 个新增世界。'],
    ['原创角色','65 名：每名含公开目标、秘密、资源、底线、关系与个人时钟。'],
    ['第三方 IP 研究','20 条：均为 REF-ONLY，附官方/权利方来源链接；不构成授权。'],
    ['原创视觉','16 张：均为 ORIG-AI-VIS，需在正式发布前人工审核。'],
    ['关系范式','11 条：支持权力、照护、秘密、身份、跨代与非人行动者的长期演化。'],
    ['优先浏览','先按 01_原创_世界 的题材、状态和优先级筛选，再用世界 ID 联动角色、视觉与关系。'],
    ['版权边界','第三方素材不用于产品、训练、广告、商店页或视觉生成提示；仅保存链接和抽象结论。']
]
for r, (k, v) in enumerate(summary, 3):
    index.cell(r,1,k).font = Font(bold=True, color='12324A')
    index.cell(r,1).fill = accent_fill
    index.cell(r,2,v).alignment = Alignment(wrap_text=True, vertical='top')
    index.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c in range(1,5): index.cell(r,c).border = Border(bottom=thin)
    index.row_dimensions[r].height = 38
index['A13'] = '使用与筛选规则'
index['A13'].font = Font(bold=True, color='FFFFFF')
index['A13'].fill = subheader_fill
index.merge_cells('A13:D13')
notes = [
    ['绿色行','本轮 v0.2 新增或孵化条目；蓝色行是既有内容。'],
    ['资产标签','ORIG-TEXT 和 ORIG-AI-VIS 才可进入内部原创资产审阅；REF-ONLY 永远不等于可商用。'],
    ['来源链接','链接用于核对和研究，不是媒体下载入口，也不构成授权文件。'],
    ['世界状态','优先用初始状态、核心冲突与长期钩子搭建模拟变量，不将 NPC 降为任务分发器。']
]
for r, (k, v) in enumerate(notes, 14):
    index.cell(r,1,k).font = Font(bold=True)
    index.cell(r,2,v).alignment = Alignment(wrap_text=True)
    index.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    index.row_dimensions[r].height = 32
for col, w in {'A':22,'B':50,'C':18,'D':18}.items(): index.column_dimensions[col].width = w
index.freeze_panes = 'A3'

world_ids_new = {f'W-{i:02d}' for i in range(7,17)}
char_ids_new = {f'C-{i:02d}' for i in range(26,66)}
ip_ids_new = {f'IP-{i:02d}' for i in range(13,21)}
vis_ids_new = {f'VIS-O-{i:02d}' for i in range(7,17)} | {'REF-02'}
rel_ids_new = {f'RLP-{i:02d}' for i in range(6,12)}

add_sheet('01_原创_世界',['世界 ID','名称','资产属性','题材 / 基调','Simulation 核心','初始状态','核心地点','阵营','主冲突','长期钩子','建议优先级','关联视觉'],worlds,[11,22,16,28,35,38,34,28,34,28,14,13],'OriginalWorlds',new_id_prefixes=world_ids_new)
add_sheet('02_原创_角色',['角色 ID','世界 ID','名称','公开身份','公开目标','秘密','资源','底线','关键关系','关联任务'],characters,[10,10,14,22,25,36,23,27,32,24],'OriginalCharacters',new_id_prefixes=char_ids_new)
ip_ws = add_sheet('03_IP_参考库',['IP ID','IP / 世界','媒介','类型','Simulation 适配度','可借鉴结构','高张力组合','地点 / 阵营 / 事件','官方来源链接','使用标签'],ips,[10,25,18,22,17,35,30,35,48,15],'IPReferences',link_cols={9},new_id_prefixes=ip_ids_new)
for r in range(2, len(ips)+2):
    ip_ws.cell(r,10).fill = warn_fill
    ip_ws.cell(r,10).font = Font(color='9C1C1C', bold=True)
add_sheet('04_视觉资产',['资产 ID','属性','名称','关联世界','文件 / 来源','规格','视觉关键词','权属标签','使用备注'],visuals,[12,14,26,13,52,20,30,18,38],'VisualAssets',new_id_prefixes=vis_ids_new)
add_sheet('05_授权矩阵',['标签','定义','商业使用指引','留档要求'],licenses,[20,48,56,38],'LicenseMatrix')
add_sheet('06_关系范式',['关系 ID','关系范式','适用组合','持续 Simulation 机制','典型演化','设计防呆'],relations,[12,24,34,43,31,42],'RelationshipPatterns',new_id_prefixes=rel_ids_new)
add_sheet('07_本轮变更',['版本','日期','类别','变更范围','内容摘要'],change_log,[12,18,16,35,62],'ChangeLog')

for sheet in wb.worksheets:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0

xlsx = OUT / 'World_Character_素材库_v0.2.xlsx'
wb.save(xlsx)
print(json.dumps(payload['metadata']['counts'], ensure_ascii=False))
print(xlsx)
print(OUT / 'world_character_library_data_v0.2.json')
