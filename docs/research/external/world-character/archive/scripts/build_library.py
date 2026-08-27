from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
import json

ROOT = Path('/home/ubuntu/world_character_library')
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)

worlds = [
    ['W-01','雾港九局','完全原创','近未来都市悬疑；温情群像','记忆雾使城市安全与身份完整性互相冲突','雾浓度 62；市民信任 48；港区失业 31；申诉积压 117','雾堤旧港；九局大楼；夜航诊所；潮汐档案馆；高线；净雾塔','市政记忆局；港区互助会；息海；回声室；盐灯','安全 vs 真相；疗愈 vs 同意；繁荣 vs 尊严','第十局、明天的尸体、净雾塔检修窗','高','VIS-O-01'],
    ['W-02','白昼失窃案','完全原创','太阳朋克群岛；生态政治；航海冒险','每日失去的阳光与日照债务重新分配岛屿命运','失昼 7分14秒；日照债务 41亿光分；珊瑚农场健康 53；苏醒度 9','向日港；桑叶环；赤道镜阵；雾藻修道院；照寂；风信塔','七岛光盟；借光银行；海藻修会；无灯渔团；深箔家','饥荒救援 vs 生态修复；债务 vs 能源公地','发光藻种、刻度贝、海底校准','高','VIS-O-02'],
    ['W-03','纸灯镇的最后一周','完全原创','现代校园；超自然悬疑；成长喜剧','毕业周循环，察觉者越多，镇外世界越褪色','循环第19次；褪色值 37；记忆者 6；未发表稿 12','纸灯高中；旧戏院；河堤夜市；天文台；报刊库；舞会礼堂','校刊社；戏剧社；校队；家长会；镇史协会；留校生','留住当下 vs 走向未来；私人遗憾 vs 集体时间','第零期、成人共读、舞会最后一首歌','高','VIS-O-03'],
    ['W-04','冰原收容所 17','完全原创','末日生存；治理模拟；家园建造','资源配给制度与集体合法性会互相消耗','人口 1023；食物 43天；热能 58；车体裂缝 11；支持率 52','母舱；温室舱；旧铁路隧道；冰裂峡；气象站；盲眼地热井','理事会；维修班；温室协作社；流动学校；第三炉；外来车队','效率 vs 不可量化的人；公开治理 vs 危机独裁','白票、地热井、外来车队','高','VIS-O-04'],
    ['W-05','红绸王朝的夜班','完全原创','历史幻想；宫廷谍报；治愈日常','梦诏可改写现实，但会造成失忠症与政令矛盾','梦诏一致性 59；粮仓 63；失忠症 23；皇帝清醒 17分/日','朱雀内廷；长梦库；御膳房夜市；织云坊；雨廊学宫；兵符渡口','夜班六司；太后内廷；边军使团；金针；无名簿','王朝稳定 vs 清醒真相；忠诚作为德性 vs 技术','醒龙灯、三道赈灾梦诏、国梦','高','VIS-O-05'],
    ['W-06','星渊合唱团','完全原创','太空歌剧；演艺群像；异星外交','音乐换取航线权，但会永久磨损团员情感','燃料 46；债务 72；和声共鸣 15；情感磨损 8','不谢幕号；零重力剧场；索因；歌剧卫星；莫拉；星渊边界','合唱团；艺人公会；航线财团；静音者；倒带','艺术自由 vs 外交商品；家庭 vs 个人成名','返场录音带、无声演出、情感备份库','高','VIS-O-06'],
    ['W-07','鲸落便利店','完全原创种子','温馨生活；幽灵经济；都市奇幻','午夜便利店随鲸骨街漂移，顾客用未说出口的话支付','待孵化','鲸骨街；夜班便利店；失物柜；潮退站','店员联盟；幽灵配送站；街区租赁会','照顾他人 vs 自我表达','漂移街区、未说出口的货币','中','待生成'],
    ['W-08','祈雨列车','完全原创种子','民俗奇幻；旅行；公路叙事','车票以放下的一段执念支付；每站有地方神干预','待孵化','祈雨列车；干涸站台；神社货场；雨云顶棚','列车局；地方神；气象商会；旅客互助会','离别 vs 执念；现代交通 vs 地方信仰','雨票、车长印、被扣留的车厢','中','待生成'],
    ['W-09','第七次小行星婚礼','完全原创种子','浪漫喜剧；外交；时间循环','外交事故让婚礼重办，每次新娘选择不同盟友','待孵化','小行星宴会厅；观景舱；外交通道','婚礼委员会；外交团；反殖民者；娱乐媒体','真爱 vs 政治联姻；真实身份 vs 公共叙事','七份誓言、损坏的星图','中','待生成'],
    ['W-10','三重城墙','完全原创种子','战争；城市建设；阶层政治','每十年外扩一圈城墙，墙外居民自动失去公民权','待孵化','内城；旧墙市场；新墙工地；墙外诊所','建设署；城防军；墙外联合会；档案局','安全扩张 vs 公民权；空间资源 vs 身份','缺口地图、十年钟、石料配额','中','待生成']
]

characters = [
    ['C-01','W-01','林雁回','九局记忆申诉官','替普通人取回误删人生','每周收到未来自己寄来的记忆，预示她会杀同事','申诉档案','不强迫他人恢复记忆','苏缄：信任/隐瞒；陶映：救命债','找出明天的尸体'],
    ['C-02','W-01','苏缄','港警重案警探','追查无动机自首','删去妹妹死亡当天记忆，妹妹可能在回声室','警权与旧港线人','不以无辜者作诱饵','林雁回：搭档；闻钧：旧怨','找回妹妹线索'],
    ['C-03','W-01','陶映','夜航诊所记忆修复师','守住港区病人','诊所兼作盐灯安全屋，制作疼痛地图','社区声望与医疗设备','不交出患者数据','林雁回：伦理争执；阿砾：保护','转移36名患者'],
    ['C-04','W-01','闻钧','息海财团公关总监','维持净雾塔续约','净雾训练预测治理模型，父亲是试验受害者','媒体与预算','不牺牲儿童','苏缄：旧怨/合作','泄露或掩盖模型'],
    ['C-05','W-01','阿砾','高线送信少年','让姐姐离开雾港','不受记忆雾影响，涂鸦会唤回记忆','全城送信网络','不让姐姐再被利用','陶映：被保护；各阵营争夺','完成全城壁画'],
    ['C-06','W-02','岑光','少年气象师','测出失昼规律','视网膜植入旧文明观测协议','风信塔访问权','不牺牲一座岛救另一座','伊塔：技术伦理冲突；母亲：债务','护送藻种'],
    ['C-07','W-02','伊塔·泊','无灯渔团船长','让底层获得食物与光','错误救援致83人失明，有关镜阵航图','船队与黑市','不让孩子参加武装行动','弥野：旧情/政治敌手','关闭或保住镜阵'],
    ['C-08','W-02','弥野','七岛光盟执政官','推行日照基本配给','挪用外交资金维持镜阵','合法权力与演讲力','不允许政治暗杀','伊塔：旧情；岑光：证据依赖','议会投票'],
    ['C-09','W-02','霖婆','海藻修会修士','保护发光藻生态','失昼是被延迟的世界校准','种子库与修会网络','不毁活体生态','各岛：精神影响','让校准完成或延迟'],
    ['C-10','W-03','宁栀','校刊主编','出版最终刊','匿名专栏是循环锚，害怕被遗忘','校刊档案与组织力','不公开匿名故事','贺川：讽刺到信任；程惟：共谋','撕毁或改写锚文'],
    ['C-11','W-03','贺川','校队王牌','赢最后一场比赛','知道镇外褪色，想偷更多陪母亲的时间','人气与体育馆钥匙','不用他人未来换家人','宁栀：信任危机；程惟：秘密交易','决定结束循环'],
    ['C-12','W-03','程惟','戏剧社灯光师','守住将消失的家族','来自旧镇、可能为留校生后代','戏院与镇史残卷','不能伤害祖母','罗青：变量依赖；宁栀：共同调查','越过镇界'],
    ['C-13','W-03','罗青','转学生、夜市摊主之女','攒钱离开小镇','不记得循环却留下梦游笔记','夜市消息网与直觉','不当实验对象','众人：不可控变量','留下下一轮信息'],
    ['C-14','W-04','顾疏','总维修师','暴风前修好主履带','评分系统由她训练且低估照护劳动','工程班与检修权','不让系统决定生死','柯晚：互信/争论；阿壹：照顾','检修与揭露算法'],
    ['C-15','W-04','柯晚','临时理事长','维持公开表决','伪造健康评估，外来车队有失散弟弟','法规与广播','不取消申诉权','顾疏：治理冲突；弟弟：亲属压力','外来车队合并'],
    ['C-16','W-04','马提亚','温室农艺师','将温室变为公地','藏有高产但毁土菌种','食物生产与工人支持','不拿儿童配给谈判','第三炉：旧债','菌种取舍'],
    ['C-17','W-04','阿壹','流动学校学生','成为记录员','能听到预测系统残响','学生网络与旧图书馆','不让朋友被赶下车','系统：人性接口；顾疏：被照顾','解读系统'],
    ['C-18','W-05','谢折枝','梦诏校勘官','为姐姐翻案','笔迹与皇帝梦中另一个自己一致','长梦库索引','不篡改亲人记忆','裴观澜：互助；太后：操控','校正赈灾梦诏'],
    ['C-19','W-05','裴观澜','夜班御膳掌事','守住夜班厨房','前边军密探，知道战报被操纵','后勤与情报','不以饥饿统治','司秋：旧情；谢折枝：同盟','寿宴试探'],
    ['C-20','W-05','司秋','边军女使者','取得真实兵符','感染失忠症，将友人误作妹妹或敌人','武艺与边军名义','不让部下替病送命','裴观澜：旧情；谢折枝：投射','公开或隐瞒疾病'],
    ['C-21','W-05','太后闻知微','白日代政者','防止王朝碎裂','维持梦政是为掩盖旧战争罪','内廷与财税','不让孙辈作人质','谢折枝：利用/潜在盟友','叫醒皇帝'],
    ['C-22','W-06','弥砂','团长兼主唱','让团员摆脱债务','每次演唱遗忘对母亲的一次怨恨','合同与舞台魅力','不签卖身演出协议','阙郎：创作搭档/权力不等','自由航线'],
    ['C-23','W-06','阙郎','作曲与引擎师','解析星渊和声','把情感备份给飞船，备份人格正代替他','维修与音乐算法','不制造有感知的奴隶','备份人格：同一性冲突','销毁或保留备份'],
    ['C-24','W-06','宝鹤','少年鼓手与直播红人','证明自己不是团里的孩子','能听懂静音者，发现他们不想被治愈','粉丝与传播力','不把文化当噱头','静音者：跨文化友谊；弥砂：监护','无声演出'],
    ['C-25','W-06','海琳','巡演经纪、前财团谈判官','偿还造成的殖民债务','为倒带盗播组织筹资难民站','合同与航线','不出卖团员身份','弥砂：理念冲突；倒带：秘密盟友','公开或藏匿倒带身份']
]

ips = [
    ['IP-01','鬼灭之刃','漫画/动画/电影','暗黑奇幻；小队任务','A','任务小队；身份转换；伤势与家族秘密','炭治郎/祢豆子；主角小队；柱/新队员','任务据点；锻刀工坊；偏远村落；人/鬼边界','https://demonslayer-anime.com/risshihen/','REF-ONLY'],
    ['IP-02','降世神通：最后的气宗','动画/电影/衍生作品','元素奇幻；旅行政治','A','区域文化与能力、国际关系耦合','安昂/祖寇；卡塔拉/索卡；导师/主角','角色/时间线/旅行区域','https://www.avatarstudiosofficial.com/','REF-ONLY'],
    ['IP-03','星露谷物语','游戏','温馨生活；社区','A','居民日程；季节；友谊与家庭','新居民/镇长；邻居/社区组织者','农场；鹈鹕镇；节日；洞穴','https://www.stardewvalley.net/','REF-ONLY'],
    ['IP-04','黑帝斯','游戏','神话；循环；家庭','A','失败后关系推进；基地重返','主角/父亲；主角/导师；旧恋人','基地；逃离路线；每轮返回','https://www.supergiantgames.com/games/hades/','REF-ONLY'],
    ['IP-05','沙丘：觉醒/沙丘宇宙','小说/电影/游戏','生态政治；沙漠生存','A','环境变量；领地建设；社会组织','生存专家/家族使者；向导/建设者','动态厄拉科斯；基地；资源线路','https://duneawakening.com/','REF-ONLY'],
    ['IP-06','辐射','游戏/剧集','末日；阵营；黑色幽默','A','区域演化；势力冲突；季节更新','幸存者/社区；理想者/民兵；商人/流浪者','废土区域；据点；入侵事件','https://fallout.bethesda.net/en','REF-ONLY'],
    ['IP-07','巫师','小说/游戏/剧集','职业冒险；地方政治','A','委托层级；道德代价；职业中立破裂','职业者/委托人；怪物/社区','地区任务；谣言；隐藏受益人','https://www.thewitcher.com/sg/en/witcher4','REF-ONLY'],
    ['IP-08','苍穹浩瀚','小说/电视剧','硬科幻；阶层；航线政治','A','小队与多方政治网络；资源约束','船员家庭/国家任务；内行星/带民','舰船；港口；航窗；环门','https://www.primevideo.com/detail/0PWWY8BZB8COGDELCHYLRXOCR5','REF-ONLY'],
    ['IP-09','最终幻想 XIV','MMO','持续世界；城邦；公共事件','A','城邦/种族；公共威胁；重复危机的历史感','冒险者/城邦代表；不同族群伙伴','Locations；Races；Threats；公共事件','https://na.finalfantasyxiv.com/a_realm_reborn/world/#!/lore','REF-ONLY'],
    ['IP-10','博德之门 3','游戏','队伍关系；选择后果','A','伙伴秘密；债务；个人任务；离队/转向','伙伴/旧主；信仰/失忆者；倒计时角色/团队','营地；旅行；同伴任务','https://baldursgate3.game/','REF-ONLY'],
    ['IP-11','火焰之纹章：风花雪月','游戏','校园；家系；战争','A','日常关系到战争阵营的时间相位','同学/未来对手；师生；领袖群','学院；三领地；战后重建','https://www.nintendo.com/us/store/products/fire-emblem-three-houses-switch/','REF-ONLY'],
    ['IP-12','权力的游戏','小说/电视剧','家族政治；战争','A','血缘/婚盟/领地/外部季节时钟','继承人/非婚生亲属；婚盟；边境/首都','领地；宴席；边境；冬季威胁','https://www.hbomax.com/show/4f6b4985-2dc9-4ab6-ac79-d60f0860b0ac','REF-ONLY']
]

visuals = [
    ['VIS-O-01','原创','雾港九局封面','W-01','/visuals/original_w01_fogharbor_cover.png','16:9；2560×1440','青灰/琥珀；雨港；记忆雾','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-02','原创','白昼失窃案封面','W-02','/visuals/original_w02_daylight_heist_cover.png','16:9；2560×1440','太阳朋克；镜阵；群岛农场','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-03','原创','纸灯镇场景','W-03','/visuals/original_w03_paper_lantern_town_scene.png','16:9；2560×1440','纸灯夜市；群像；时钟','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-04','原创','冰原收容所场景','W-04','/visuals/original_w04_ice_shelter_scene.png','16:9；2560×1440','极光；履带车；温室剖面','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-05','原创','红绸王朝场景','W-05','/visuals/original_w05_red_silk_dynasty_scene.png','16:9；2560×1440','朱红；雨廊；梦诏灯','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['VIS-O-06','原创','星渊合唱团封面','W-06','/visuals/original_w06_star_choir_cover.png','16:9；2560×1440','零重力；星云；舞台光','ORIG-AI-VIS','内部原创生成；发布前人工审核'],
    ['REF-01','第三方','官方世界/角色/预告片入口','IP-01 至 IP-12','见 IP_参考库工作表中的来源链接','网页/视频','仅用于镜头、色彩、关系表达研究','REF-ONLY','不可下载、嵌入、训练或商用；如需使用须另取权利人许可']
]

licenses = [
    ['ORIG-TEXT','本项目原创文本世界观/角色/关系/事件','可在完成内部权属登记后进入商业产品','记录版本、作者/生成记录、编辑历史'],
    ['ORIG-AI-VIS','为本项目生成且不含第三方角色/标识/复刻指令的图像','建议人工审核、核查生成平台条款和商标/肖像风险后使用','保留提示词、生成日期、源文件、审阅记录'],
    ['ORIG-COMMISSION','根据书面委托产出的原创视觉','依合同；建议取得商业使用与衍生权','保存合同、交付清单、付款与权属条款'],
    ['REF-ONLY','第三方作品、官方媒体、搜索到的图片/视频、粉丝内容','不可直接用于产品、广告、商店页、训练集或衍生物','仅保存链接、观察笔记与抽象化结论'],
    ['LICENSE-REQUIRED','有意使用第三方可识别角色/视觉/音视频','取得书面许可后才可使用，范围严格按合同','保存授权方、地域、媒介、期限、署名义务']
]

relation_patterns = [
    ['RLP-01','相互需要的对立同盟','警探/公关；船长/执政官','共同外敌 + 利益相反 + 资源互赖','合作 → 背叛 → 再谈判','双方各保留不可放弃的善意底线'],
    ['RLP-02','保护者与不可控变量','医师/免疫少年；团员/跨文化儿童','保护会衍生控制，变量会撬动权力','庇护 → 争夺 → 夺回行动权','未成年角色需清晰安全边界，不承担成人浪漫化功能'],
    ['RLP-03','旧情人与公共职责','船长/执政官；密探/使者','私人史持续影响公共决策','隐瞒 → 利用 → 坦白/分道扬镳','复合不应是唯一好结局'],
    ['RLP-04','人与制度性复制物','工程师/预测系统；作曲家/备份人格','把效率、同意、身份转成日常互动','依赖 → 不信任 → 共治或断绝','复制物要有知识边界、偏好与责任'],
    ['RLP-05','秘密互保的弱关系','摊主/主编；御膳掌事/校勘官','小秘密可以长成信任也能成为背刺','保密 → 试探 → 揭露/互保','每个秘密至少提供坦白、揭穿、误解三条线']
]

summary = [
    ['素材库名称','World / Character 内容与视觉素材库 v0.1'],
    ['创建日期','2026-08-25（GMT+8）'],
    ['原创世界','10 条：6 个旗舰世界 + 4 个孵化种子'],
    ['原创角色','25 名：每名包含目标、秘密、资源、底线、关系与任务'],
    ['第三方 IP 参考','12 条：附官方来源链接与 REF-ONLY 标注'],
    ['原创视觉','6 张：均为 ORIG-AI-VIS，需发布前审阅'],
    ['推荐浏览路径','先筛选原创_世界表的优先级和题材，再按世界 ID 联动原创_角色表、视觉资产表及关系范式表。'],
    ['重要边界','第三方素材不等于可商用；本库只将其作为结构/镜头/关系研究。']
]

payload = {
    'original_worlds': worlds,
    'original_characters': characters,
    'ip_references': ips,
    'visual_assets': visuals,
    'license_matrix': licenses,
    'relationship_patterns': relation_patterns
}
(OUT / 'world_character_library_data.json').write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')

wb = Workbook()
wb.remove(wb.active)

header_fill = PatternFill('solid', fgColor='12324A')
subheader_fill = PatternFill('solid', fgColor='1F567D')
accent_fill = PatternFill('solid', fgColor='E7F1F7')
warn_fill = PatternFill('solid', fgColor='FBE9E7')
thin = Side(style='thin', color='C6D4DF')


def add_sheet(name, headers, rows, widths=None, table_name=None, source_link_cols=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = 'A2'
    for col_idx, value in enumerate(headers, 1):
        cell = ws.cell(1, col_idx, value)
        cell.fill = header_fill
        cell.font = Font(color='FFFFFF', bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = Border(bottom=thin)
    for row_idx, row in enumerate(rows, 2):
        for col_idx, value in enumerate(row, 1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.border = Border(bottom=thin)
            if row_idx % 2 == 0:
                cell.fill = accent_fill
            if source_link_cols and col_idx in source_link_cols and isinstance(value, str) and value.startswith('http'):
                cell.hyperlink = value
                cell.style = 'Hyperlink'
        ws.row_dimensions[row_idx].height = 44
    ws.row_dimensions[1].height = 32
    for col_idx in range(1, len(headers) + 1):
        width = (widths[col_idx - 1] if widths else 18)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    if rows:
        tab = Table(displayName=table_name or f'Table{len(wb.sheetnames)}', ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
        tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showFirstColumn=False, showLastColumn=False, showRowStripes=False, showColumnStripes=False)
        ws.add_table(tab)
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(2, len(rows)+1)}'
    return ws

# Index / readme
ws = wb.create_sheet('00_索引')
ws.sheet_view.showGridLines = False
ws.merge_cells('A1:D1')
ws['A1'] = 'World / Character 内容与视觉素材库'
ws['A1'].font = Font(size=20, bold=True, color='FFFFFF')
ws['A1'].fill = header_fill
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 38
for r, (key, value) in enumerate(summary, 3):
    ws.cell(r, 1, key).font = Font(bold=True, color='12324A')
    ws.cell(r, 1).fill = accent_fill
    ws.cell(r, 2, value).alignment = Alignment(wrap_text=True, vertical='top')
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    for c in range(1,5):
        ws.cell(r,c).border = Border(bottom=thin)
    ws.row_dimensions[r].height = 34
ws['A13'] = '筛选字段说明'
ws['A13'].font = Font(bold=True, color='FFFFFF')
ws['A13'].fill = subheader_fill
ws.merge_cells('A13:D13')
filter_notes = [
    ['优先级','仅代表建议先做原型验证的顺序，不代表最终产品决策。'],
    ['题材/类型','可以跨表筛选并组合；一个世界可支撑多个题材标签。'],
    ['资产标签','ORIG-TEXT、ORIG-AI-VIS、REF-ONLY 等决定可用方式。'],
    ['来源链接','所有第三方 IP 都指向官方/权利方入口，链接只是研究入口，不是授权证明。']
]
for r, row in enumerate(filter_notes, 14):
    ws.cell(r,1,row[0]).font = Font(bold=True)
    ws.cell(r,2,row[1]).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    ws.row_dimensions[r].height=28
ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 46
ws.column_dimensions['C'].width = 18
ws.column_dimensions['D'].width = 18
ws.freeze_panes = 'A3'

add_sheet('01_原创_世界', ['世界 ID','名称','资产属性','题材 / 基调','Simulation 核心','初始状态','核心地点','阵营','主冲突','长期钩子','建议优先级','关联视觉'], worlds, [11,22,16,28,35,38,34,28,34,28,14,13], 'OriginalWorlds')
char_ws = add_sheet('02_原创_角色', ['角色 ID','世界 ID','名称','公开身份','公开目标','秘密','资源','底线','关键关系','关联任务'], characters, [10,10,14,22,25,36,23,27,32,24], 'OriginalCharacters')
add_sheet('03_IP_参考库', ['IP ID','IP / 世界','媒介','类型','Simulation 适配度','可借鉴结构','高张力组合','地点 / 阵营 / 事件','官方来源链接','使用标签'], ips, [10,22,18,22,17,35,30,35,48,15], 'IPReferences', {9})
vis_ws = add_sheet('04_视觉资产', ['资产 ID','属性','名称','关联世界','文件 / 来源','规格','视觉关键词','权属标签','使用备注'], visuals, [12,14,24,13,52,20,30,18,37], 'VisualAssets')
add_sheet('05_授权矩阵', ['标签','定义','商业使用指引','留档要求'], licenses, [20,48,56,38], 'LicenseMatrix')
add_sheet('06_关系范式', ['关系 ID','关系范式','适用组合','持续 Simulation 机制','典型演化','设计防呆'], relation_patterns, [12,24,34,43,31,42], 'RelationshipPatterns')

# Conditional formatting and link to external reference sheet
for row in range(2, len(worlds) + 2):
    if ws.title != '01_原创_世界':
        pass
world_ws = wb['01_原创_世界']
world_ws.conditional_formatting.add(f'K2:K{len(worlds)+1}', ColorScaleRule(start_type='min', start_color='FCE4D6', mid_type='percentile', mid_value=50, mid_color='FFF2CC', end_type='max', end_color='C6E0B4'))
for sheet_name in ['01_原创_世界','02_原创_角色','03_IP_参考库','04_视觉资产','05_授权矩阵','06_关系范式']:
    wb[sheet_name].sheet_properties.pageSetUpPr.fitToPage = True
    wb[sheet_name].page_setup.fitToWidth = 1
    wb[sheet_name].page_setup.fitToHeight = 0

# IP license column visual warning
for row in range(2, len(ips) + 2):
    wb['03_IP_参考库'].cell(row, 10).fill = warn_fill
    wb['03_IP_参考库'].cell(row, 10).font = Font(color='9C1C1C', bold=True)

xlsx = OUT / 'World_Character_素材库_v0.1.xlsx'
wb.save(xlsx)
print(xlsx)
print(OUT / 'world_character_library_data.json')
