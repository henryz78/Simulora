import json
from pathlib import Path
from copy import deepcopy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

ROOT = Path('/home/ubuntu/world_character_library')
OUT = ROOT / 'deliverables'
with (OUT / 'world_character_library_data_v0.4.json').open('r', encoding='utf-8') as f:
    base = json.load(f)

worlds = deepcopy(base['original_worlds'])
characters = deepcopy(base['original_characters'])
ips = deepcopy(base['ip_references'])
visuals = deepcopy(base['visual_assets'])
licenses = deepcopy(base['license_matrix'])
relations = deepcopy(base['relationship_patterns'])
change_log = deepcopy(base['change_log'])

worlds.extend([
['W-39','逆潮物流网','完全原创','物流与供应链；港口；区域公共服务','台风季提前，自动排序压低药品、冷链食品和救灾物资；港口与内陆社区必须重新定义谁有资格给货物排序','冷链完整度 62；药品延误 84 箱；船期可信度 43；司机疲劳 71；社区库存 9 天','潮汐集装箱码头；夜班调度塔；河驳船换装站；冷链仓；司机休息院；山地合作社','港务自动化联盟；司机合作社；药品配送协会；出口商理事会；内河船员工会；社区采购联盟','效率 vs 生存优先；商业机密 vs 调度透明；司机安全 vs 保供压力','失联驳船；人工分流；红色需求图；赔付优先港/共同排程网/黑市走廊','高','VQ-39'],
['W-40','鹿角岛师范学院','完全原创','教育机构；师范培养；岛屿社区','偏远学校支援基金与毕业生去向绑定，学院必须决定教育是留住人还是替城市筛选人','实习岗位 83；乡村教师空缺 29；学生债务 46；家长信任 57；渡船天气窗 5 天','海风讲堂；模拟教室；渡船码头；山村寄宿学校；教师宿舍；儿童广播站','学院理事会；学生教师工会；乡村校长联盟；家长互助会；在线教育平台；船工会','标准化成绩 vs 本地知识；职业机会 vs 长期承诺；儿童表达权 vs 成人代言','儿童广播；基金复核；暴风寄宿；派遣工厂/岛校共同体/儿童共治学园','高','VQ-40'],
['W-41','白昼试验站','完全原创','科研团队；极端环境；开放科学','极昼仅余 46 天，冰下微生物可能储热；采样过度会破坏食物链，排他专利又在逼近','样本稳定度 49；燃料 38 天；研究信任 52；发表窗口 19 天；许可 1 次','冰下观测井；温室实验舱；样本冷库；共同数据室；海岸渔营；直升机坪','驻站科学组；生态知识委员会；产业实验室；开放数据联盟；飞行队；伦理观察员','发现速度 vs 生态节制；开放科学 vs 研究劳动；地方同意 vs 远程资助','钻井裂缝；共同署名；极昼节；专利前哨/共管研究站/生态修复营','高','VQ-41'],
['W-42','栖灯消防分队','完全原创','公共安全；灾害响应；邻里防灾','老旧街区连续小火警让保险、租住、维修和恐慌相互放大；分队要在执法与居住稳定间寻找长期预防','响应 8分42秒；水压 54；装备 6；居民信任 36；演练参与 19%','栖灯消防站；屋顶水箱；社区厨房；教育车；保险申诉亭；安置馆；检修巷','消防分队；租户委员会；房东联合会；保险风险组；维修工；青年志愿队；监察处','执法 vs 居住稳定；通报 vs 恐慌；保险模型 vs 街区修复','隐患图；水箱抢修；保险听证；高风险驱离区/共防街区/警报疲劳区','高','VQ-42'],
['W-43','白瓷公社餐桌','完全原创','农业；食品系统；餐饮；合作社','谷物病害威胁农户、学校厨房、食堂和餐馆共同食品网；廉价餐与多样性种植不能再被分开处理','主粮健康 41；学校餐覆盖 78%；厨师工时 63；现金 26 天；种子多样性 33','社区谷仓；中央厨房；餐桌市场；菌种实验棚；罐藏工坊；家庭农场；食物银行','小农合作组；学校餐联盟；餐馆采购会；食物银行；农技商社；儿童食谱委员会','均价餐食 vs 农户生计；安全 vs 浪费；志愿热情 vs 组织剥削','病害隔离；儿童盲评；丰收宴取消；单一供应餐桌/多样性食物网/应急厨房','高','VQ-43'],
['W-44','边界交换站','完全原创','跨境社区；铁路口岸；移民服务','数字通行证排除没有稳定网络与地址的人；夜车站内的翻译、医疗、汇款和家庭重聚成为替代通行实践','夜车座位 240；证件错误率 18%；翻译排队 4小时；汇款延迟 9天；信任 45','双语站厅；夜车月台；翻译台；跨境诊所；包裹柜；边市餐馆；旧护照档案室','铁路运营组；边境服务处；跨境家庭联盟；商贩协会；翻译互助会；通行证承包商','边境安全 vs 家庭连续；数字效率 vs 可达性；申诉公开 vs 证人隐私','夜车满员；翻译罢工；档案听证；数字门槛站/人际通行站/跨境共同体枢纽','高','VQ-44'],
['W-45','合页百货公司','完全原创','家族企业；零售转型；多代家庭','百年百货面对关店或共享商业转型；家族、员工、线上团队和街区商户争夺的不只是股份，还有家业定义权','现金流 28 天；家族投票 3:3；员工留任 64；线上订单增长 38；空置层 47','礼品厅；旧电梯；地下仓；顶层食堂；档案间；直播库房；修鞋摊','家族股东；老员工会；线上零售组；街区商户联盟；地产收购方；顾客回忆社','传承 vs 转型；家族控制 vs 员工所有；线上便利 vs 街区关系','遗嘱复核；停电日；员工股权案；封闭百货/共营商场/街区服务合作社','高','VQ-45'],
['W-46','余音美术馆','完全原创','艺术文化机构；修复伦理；公共记忆','旧工业城口述史录音能挽救美术馆财政，也要求机构承认撤回、沉默、纠正和归还权','修复预算 44；馆员疲劳 58；来源同意 51；游客预测 73；归还请求 6','声音修复室；临展大厅；开放档案桌；旧厂外展车；儿童工坊；保管库；河畔放映墙','策展组；讲述者联盟；修复师会；赞助委员会；青年导览团；旧厂居民会；展览供应商','保存 vs 撤回；体验 vs 安全；客流 vs 慢修复；机构权威 vs 社区共管','音频撤回；赞助条件；归还仪式；品牌沉浸馆/共管记忆馆/流动归还站','高','VQ-46'],
['W-47','环时堤岸局','完全原创','长期基础设施；地方政治；代际治理','海平面工程需 80 年完成，每五年允许未来席位参与预算；迁居、湿地与债券让当下和后代的利益冲突可见','一期完工 23%；迁居户 1208；未来席位信任 28；预算缺口 19；湿地恢复 36','堤岸局大厅；未来席位厅；迁居社区；潮沟工地；档案塔；儿童模型工坊；浮动菜场','堤岸工程局；迁居家庭联盟；未来代理团；湿地保育会；商户会；债券方；青年议会','当下住房 vs 远期安全；硬堤 vs 可逆设计；未来代表 vs 当前民主','五年听证；迁居抽签；湿地修复；硬堤开发城/可逆湿地城/迁居正义城','高','VQ-47'],
['W-48','镜界迁居署','完全原创','多世界社会；迁徙治理；跨物种公共服务','七个世界的迁居者来到中转城，住房、语言、物种需求、家庭分离与通道关闭都不能以单一世界为默认标准','通道稳定 48；临时住房 73；跨界家庭 86；翻译覆盖 61；协议 3/7','镜界抵达厅；七语服务台；适应性住房区；跨物种诊所；观测台；重聚花园；检疫库','迁居署；七界互助联盟；通道工程团；原居民会；跨物种照护组；财产申诉庭；关闭通道派','接纳 vs 承载；同化 vs 多世界尊严；返回权 vs 安全关闭','通道闭合；物品禁令；世界节；紧急中转城/多世界共居城/七界联邦港','高','VQ-48']
])

characters.extend([
['C-154','W-39','温泊','夜班港口调度员','让药品先出港','亲人等待延误透析液，曾手动改排序','泊位表；调度台；船员信任','不伪造危险品记录','宋庾：互用/互查；孟凡：调度冲突','每晚 02:00 重排滞港船，台风前锁定冷链时隙'],
['C-155','W-39','宋庾','司机合作社代表','争取安全工时','掌握未报备山路，曾让疲劳司机走过','司机网；路线；休息院','不强迫无证司机接单','温泊：共同保供；杨述：清单同盟','每日收车后统计疲劳，延误超两小时组织停运'],
['C-156','W-39','杨述','社区药房联盟采购员','找回透析液','曾低估内陆需求，隐瞒缺药名单','库存；诊所关系；需求表','不把病人名单给商业方','宋庾：依赖；孟凡：证据对手','每日走访缺药家庭，满三天公开红色库存'],
['C-157','W-39','孟凡','自动化维护工程师','修复系统公信力','为出口赔付协议写过优先级补丁','系统权限；审计日志；技术声望','不删除原始延误数据','温泊：互疑；杨述：追问','每晚备份日志，决定提交或藏匿补丁'],
['C-158','W-40','乔楠','学生教师','留在岛上任教','奖学金要求去城市学校服务三年','课堂设计；儿童信任；同学网','不把儿童反馈伪造成指标','卢令：师徒；柚子：共创','每周改课，实习结束前申请留岛豁免'],
['C-159','W-40','卢令','乡村校长兼导师','保住寄宿学校','早知学校将合并却未公开','校务；家长网；旧校钥匙','不以关闭威胁家长','乔楠：师徒；范航：政策冲突','放学后收集意见，基金复核前公开或隐瞒'],
['C-160','W-40','范航','基金项目经理','完成派遣指标','将留任定义为一年以压低流失','基金；岗位分配；数据','不伪造儿童安全记录','卢令：旧同学；柚子：质问','每周排实习，暴风停航时重排岗位'],
['C-161','W-40','柚子','儿童广播主理人','让大人听见学生','家人准备离岛，怕节目成告别','广播钥匙；同龄人；教师信任','不念出同学私信名','乔楠：共创；范航：追问','每周播出，听证日前发起谁教谁节目'],
['C-162','W-41','迟牧','微生物学家','证明储热机制','署名危机让他想先发预印本','样本；实验记录；声望','不篡改显微数据','阿栖：共著；禾曜：资助压力','极昼 01:00 复查样本，发表前开放或锁定数据'],
['C-163','W-41','阿栖','生态委员会联络员','保护食物链','祖父曾为早期采样带路','观察；许可；渔营信任','不替社区签不可撤回同意','迟牧：互信/警惕；薇沐：飞行依赖','每日查冰下变化，裂缝即暂停采样'],
['C-164','W-41','薇沐','后勤机师','保证安全撤离','飞行合同受产业实验室控制','飞行时段；燃料；救援','不在禁飞天起飞','阿栖：照应；禾曜：压力','每次天气窗决定带人、样本或设备离站'],
['C-165','W-41','禾曜','产业资助方代表','签下排他许可','家族靠寒地矿业起家，怕公众审查','资金；专利顾问；远程话语','不明知让人受伤','迟牧：职业诱惑；阿栖：对立','每周向总部报进展，撤离前提出或撤销条款'],
['C-166','W-42','庄火','消防分队长','缩短出警时间','曾参与错误疏散，怕重上新闻','出警权；训练；队员信任','不让居民留在未解释风险里','苏圆：执法冲突；纪拾：信息','每日晨检，红色警报时优先上门'],
['C-167','W-42','苏圆','租户委员会成员','阻止整治变驱逐','家人藏有违规电路','租户网；申诉；厨房','不用邻居秘密谈判','庄火：互疑；段慈：盟友','每晚收火险，租约威胁时组织检查'],
['C-168','W-42','段慈','街区维修工','修好水箱和断路器','使用过劣质断路器填欠款','工具；图纸；同行网','不让学徒背责任','苏圆：同盟；裴槐：压力','48 小时内修关键隐患，必要时公开材料来源'],
['C-169','W-42','裴槐','保险风险评估员','避免整片拒保','模型误判低收入区为纵火高发','风险模型；理赔权限；报告','不造假火灾原因','庄火：互赖；段慈：追查','每周提交评分，听证前校正或掩盖模型'],
['C-170','W-43','唐谷','小农合作组长','稳定主粮','家庭曾推动单一种子政策','土地；种子库；农户信任','不把病害推给学徒','梁禾：品种争论；童碗：食谱盟友','每日巡田，病害扩散时隔离或冒险收割'],
['C-171','W-43','梁禾','学校中央厨房主厨','保证学生餐','长期依赖一种廉价谷物','厨房排期；菜谱；厨师团','不减少儿童餐量遮掩预算','唐谷：供给互赖；童碗：监督','每周改菜，库存低三天启动替代菜单'],
['C-172','W-43','童碗','儿童食谱委员会代表','让孩子有选择','对某些食物过敏未告诉家人','盲评；同学网；广播','不替同学投票','梁禾：合作；司荞：压力','每周盲评，食谱日发起传统菜表决'],
['C-173','W-43','司荞','餐馆采购会代表','维持餐饮就业','签过高价独家采购意向','订单；冷链；媒体','不让农户低于成本卖粮','唐谷：商业冲突；梁禾：余料合作','每晚确认订单，丰收宴前公开或取消协议'],
['C-174','W-44','艾澄','双语站务员','让夜车准点出发','姐姐无证滞留另一侧','排程；站务人脉；双语','不用旅客信息换个人通行','乔旎：协作；程栩：冲突','每晚 21:00 核座，满员时优先家庭重组'],
['C-175','W-44','乔旎','翻译互助会成员','缩短翻译等待','曾错译医疗同意书','语言；志愿者；诊所关系','不在不理解语境代签','艾澄：同盟；程栩：监督','每日排班，医疗个案优先'],
['C-176','W-44','程栩','通行证工程师','修正错误率','母亲也被姓氏模型误判','系统权限；日志；技术团','不删受害者申诉','乔旎：监督；阿那：档案依赖','每周更新，夜车前开放或关闭人工通道'],
['C-177','W-44','阿那','旧护照档案管理员','找回家庭通行权','档案含证人地址','档案；法律；老城记忆','不以证明权利暴露安全','艾澄：信任；程栩：数据冲突','每日核档案，听证前提出最小披露方案'],
['C-178','W-45','余页','家族第三代股东','让百货转型存活','私下联系收购方','股份；品牌；会议权','不让员工从媒体得知裁员','昭叔：亲情/对立；闻笑：合作','每周看现金流，家宴投票前撤销或推进收购'],
['C-179','W-45','昭叔','老员工会代表','守住员工年资','知遗嘱员工基金条款却怕家族分裂','仓储知识；员工信任；旧钥匙','不把新员工当替代品','余页：代际拉扯；闻笑：师徒','每日巡地下仓，遗嘱日公开或保留账册'],
['C-180','W-45','闻笑','青年设计团队负责人','把空层变社区空间','项目受地产方资助','设计；青年客群；展陈','不驱赶街角摊贩','余页：合作；罗莱：谈判','试营业夜前修改空间方案'],
['C-181','W-45','罗莱','线上零售组负责人','维持转型收入','外泄过客户名单','订单；数据；直播库房','不让客户隐私成股份筹码','昭叔：经验互补；闻笑：冲突','每日清单，折扣战时补救或公开泄露'],
['C-182','W-46','何修','声音修复师','完成修复','曾修掉讲述者不愿公开的停顿','修复技术；磁带；档案','不伪造原始声音','叶兰：许可协商；戴雍：压力','每晚修一段，撤回到来即暂停'],
['C-183','W-46','叶兰','讲述者联盟成员','取得共同审核权','关键录音涉及自己工伤','讲述者网；同意记录；现场','不让沉默成默认同意','何修：互信；枝空：导师','每周试听，开幕前保留或撤回录音'],
['C-184','W-46','戴雍','策展人','挽救财政','接受附条件赞助','展厅；赞助；媒体','不公开讲述者住址','何修：压力；枝空：挑战','每日调展线，复核日前拒绝或妥协'],
['C-185','W-46','枝空','青年导览团成员','让孩子参与叙事','私复制录音给旧厂家庭','导览；儿童工坊；社群','不把创伤当打卡','叶兰：导师；戴雍：冲突','每周重录，归还日前归还或公开副本'],
['C-186','W-47','陆筝','堤岸工程师','完成一期工程','知硬堤加剧内涝','工程图；模型；工地关系','不隐瞒结构风险','南环：协作；子遥：争论','每周复核潮汐，听证前提交或压下修订'],
['C-187','W-47','南环','迁居家庭联盟代表','取得稳定住房','曾夸大受灾损失','家庭网；媒体；安置清单','不让老人因搬迁失照护','陆筝：务实；子遥：代际争论','每日走访，抽签日前核对名单'],
['C-188','W-47','子遥','未来席位代理人','保留湿地','担心只是在替成人投射','听证席；青年网；模拟工具','不替未来人给确定答案','陆筝：技术争论；冯涛：对手','每月未来工坊，听证日签或拒预算'],
['C-189','W-47','冯涛','债券方代表','确保回款','有滨水开发隐性股份','融资；法律；议会关系','不直接威胁迁居户','南环：谈判；子遥：交锋','每季催进度，缺口出现推私有化条款'],
['C-190','W-48','岑折','迁居署前线协调员','让家庭重聚','伴侣仍在将关闭的镜界','排程；前线信任；追踪','不以个案换优先','微枝：协作；祁电：冲突','每日核名单，倒计时前请求或放弃特例'],
['C-191','W-48','微枝','七界互助联盟发言者','建立多物种住房标准','来自光敏世界且隐瞒不适','互助网；多语；信誉','不把一种世界当默认尺度','岑折：互信；祁电：对立','每周查住房，标准削减时组织联署'],
['C-192','W-48','祁电','通道工程师','维持稳定','知通道是污染转移副产物','工程权限；数据；维修团','不让未准备者穿越','岑折：纠葛；栖雾：证据压力','每次波动后开放、降载或关闭通道'],
['C-193','W-48','栖雾','财产申诉庭书记员','保护物品与返回权','有原居民无法继续居住的证据','案卷；检疫；调解','不把物品当天然危险','微枝：同盟；祁电：压力','每日筛申诉，投票前公布最小风险清单']
])

ips.extend([
['IP-38','Factorio','游戏','物流；生产；能源；自动化','A','多级生产链、运输、能源和研究的耦合让局部瓶颈持续传导','资源/运输；产能/维护；效率/代价','采集；物流；生产；能源；列车','https://factorio.com/game/content','REF-ONLY'],
['IP-39','Papers, Please','游戏','边境通行；文件核验；制度后果','A','规则更新、时间压力与通行判断能让小决策持续改变收入、安全、家庭和信任','检查者/通行者；规则/例外；机构/家庭','证件；队列；规则更新；申诉','https://dukope.com/','REF-ONLY'],
['IP-40','熊家餐馆','电视剧','餐饮服务；家庭企业；团队协作','A','服务高峰、班次交接、库存、家庭债和团队质量标准会在短周期中积累成长后果','服务者/顾客；家庭/团队；质量/休息','后厨；服务；排班；复盘','https://www.fxnetworks.com/shows/the-bear','REF-ONLY']
])

visuals.extend([
['VQ-39','原创队列','逆潮物流网关键视觉','W-39','/visuals/original_w39_tide_logistics_keyart.png','16:9；仅 brief','港口；冷链；人工分流','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-40','原创队列','鹿角岛师范学院关键视觉','W-40','/visuals/original_w40_antler_teacher_college_keyart.png','16:9；仅 brief','渡船；课堂；儿童广播','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-41','原创队列','白昼试验站关键视觉','W-41','/visuals/original_w41_daylight_station_keyart.png','16:9；仅 brief','冰原；科研；共同许可','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-42','原创队列','栖灯消防分队关键视觉','W-42','/visuals/original_w42_firehouse_keyart.png','16:9；仅 brief','旧街区；水箱；共防','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-43','原创队列','白瓷公社餐桌关键视觉','W-43','/visuals/original_w43_food_commons_keyart.png','16:9；仅 brief','农场；学校厨房；食物循环','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-44','原创队列','边界交换站关键视觉','W-44','/visuals/original_w44_border_exchange_keyart.png','16:9；仅 brief','夜车；翻译；通行','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-45','原创队列','合页百货公司关键视觉','W-45','/visuals/original_w45_hinge_department_store_keyart.png','16:9；仅 brief','百货剖面；员工；街区商业','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-46','原创队列','余音美术馆关键视觉','W-46','/visuals/original_w46_afterecho_museum_keyart.png','16:9；仅 brief','修复；同意；档案归还','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-47','原创队列','环时堤岸局关键视觉','W-47','/visuals/original_w47_ringtide_seawall_keyart.png','16:9；仅 brief','湿地；堤岸；代际听证','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片'],
['VQ-48','原创队列','镜界迁居署关键视觉','W-48','/visuals/original_w48_mirror_migration_keyart.png','16:9；仅 brief','多世界；适应性住房；重聚','PENDING-QUOTA','本轮仅记录原创 brief，不生成图片']
])

relations.extend([
['RLP-23','排程者与被排程者','物流；交通；基建；服务','队列与优先级决定谁等待、谁承担风险、谁得到解释','不透明安排 → 被发现 → 共同排程/抵制/替代网络','每项排序必须能解释、申诉并允许人工例外'],
['RLP-24','传授者与地方学习者','教育；科学；农业','知识传递可被儿童、学徒和地方社区反向改变','导师主导 → 共同设计 → 地方课程/知识撤回','不把本地知识当素材，不将儿童只写成希望象征'],
['RLP-25','预防者与居住者','消防；堤岸；公共安全','风险治理会在预防、强制和迁离之间摇摆','警报 → 协商 → 共防/驱离/警报疲劳','安全必须展示住房、劳动和照护后果'],
['RLP-26','机构保管者与故事所有者','文化机构；线上社区；档案','保存、修复、审核与退出均需要持续同意','委托 → 二次解释 → 撤回/共管/归还','不把公开当默认善意，也不把沉默当可利用空白'],
['RLP-27','当前受益者与未来代理','长周期基建；多世界治理','尚未参与的人也会承受制度后果','代表 → 质疑授权 → 多时钟协商/制度崩解','未来代理应有不确定性、可替换性和退出机制'],
['RLP-28','通行者与留下者','跨境社区；迁居；多世界','流动权重塑留下者的住房、语言、工作与归属','临时协助 → 公共资源争议 → 共居/封闭/互助','避免将迁居者或原居民扁平化为单一利益集团']
])

change_log.extend([
['v0.5','2026-08-25','世界','新增 W-39 至 W-48','10 个原创世界，覆盖物流、教育、科研、公共安全、食品、跨境社区、家族零售、文化机构、长周期堤岸与多世界迁居。'],
['v0.5','2026-08-25','角色','新增 C-154 至 C-193','40 名原创角色；每名包含公开目标、隐性需求、资源、底线、关键关系、个人时钟和自主行动。'],
['v0.5','2026-08-25','关系','新增 RLP-23 至 RLP-28','新增排程、知识传递、风险预防、机构保管、未来代理和通行关系范式。'],
['v0.5','2026-08-25','IP 研究','新增 IP-38 至 IP-40','3 个 REF-ONLY 研究条目，分别补足生产物流、制度化通行与高压服务协作。'],
['v0.5','2026-08-25','视觉','新增 VQ-39 至 VQ-48','仅新增完全原创视觉 brief；按要求未尝试任何图片、封面、角色、地图或动态视觉生成，VQ-17 至 VQ-48 均保持 PENDING-QUOTA。']
])

payload = {'metadata': {'library_name':'World / Character 内容与视觉素材库','version':'v0.5','updated_at':'2026-08-25 GMT+8','counts':{'original_worlds':len(worlds),'original_characters':len(characters),'ip_references':len(ips),'visual_assets':len(visuals),'relationship_patterns':len(relations)},'license_notice':'第三方 IP 及其官方媒体均为 REF-ONLY；链接不构成许可。原创生成/结构资产需在发布前人工审核。','visual_generation_note':'按 v0.5 任务要求，本轮未尝试图像生成；VQ-17 至 VQ-48 只保存原创可执行 brief，全部保持 PENDING-QUOTA。'},'original_worlds':worlds,'original_characters':characters,'ip_references':ips,'visual_assets':visuals,'license_matrix':licenses,'relationship_patterns':relations,'change_log':change_log}
(OUT / 'world_character_library_data_v0.5.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT / 'world_character_library_data.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')

wb=Workbook(); wb.remove(wb.active)
header_fill=PatternFill('solid',fgColor='12324A'); subheader_fill=PatternFill('solid',fgColor='1F567D'); accent_fill=PatternFill('solid',fgColor='E7F1F7'); warn_fill=PatternFill('solid',fgColor='FBE9E7'); new_fill=PatternFill('solid',fgColor='E8F5E9'); pending_fill=PatternFill('solid',fgColor='FFF4CC'); thin=Side(style='thin',color='C6D4DF')
def add_sheet(name,headers,rows,widths,table_name,link_cols=None,new_ids=None,pending_ids=None):
    ws=wb.create_sheet(name); ws.sheet_view.showGridLines=False; ws.freeze_panes='A2'
    for c,value in enumerate(headers,1):
        cell=ws.cell(1,c,value); cell.fill=header_fill; cell.font=Font(color='FFFFFF',bold=True); cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); cell.border=Border(bottom=thin)
    for r,row in enumerate(rows,2):
        item_id=str(row[0])
        for c,value in enumerate(row,1):
            cell=ws.cell(r,c,value); cell.alignment=Alignment(vertical='top',wrap_text=True); cell.border=Border(bottom=thin)
            if pending_ids and item_id in pending_ids: cell.fill=pending_fill
            elif new_ids and item_id in new_ids: cell.fill=new_fill
            elif r%2==0: cell.fill=accent_fill
            if link_cols and c in link_cols and isinstance(value,str) and value.startswith('http'): cell.hyperlink=value; cell.style='Hyperlink'
        ws.row_dimensions[r].height=52
    ws.row_dimensions[1].height=34
    for c,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(c)].width=w
    if rows:
        tab=Table(displayName=table_name,ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}'); tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=False,showColumnStripes=False); ws.add_table(tab)
    return ws
index=wb.create_sheet('00_索引'); index.sheet_view.showGridLines=False; index.merge_cells('A1:D1'); index['A1']='World / Character 内容与视觉素材库 v0.5'; index['A1'].font=Font(size=20,bold=True,color='FFFFFF'); index['A1'].fill=header_fill; index['A1'].alignment=Alignment(horizontal='center'); index.row_dimensions[1].height=40
summary=[['版本','v0.5 增量扩充；以 v0.4 为基线，未覆盖或重写已有资产。'],['原创世界',f'{len(worlds)} 个：新增 W-39 至 W-48，补足物流、教育、科研、灾害、食品、跨境、家族企业、文化、代际工程与多世界公共服务。'],['原创角色',f'{len(characters)} 名：新增 C-154 至 C-193；每名具备目标、秘密、资源、底线、关系、个人时钟和自主行动。'],['第三方 IP 研究',f'{len(ips)} 条：新增 IP-38 至 IP-40；全部为 REF-ONLY。'],['视觉资产',f'{len(visuals)} 条：新增 10 组原创视觉 brief；严格按要求不生成图像，VQ-17 至 VQ-48 全部为 PENDING-QUOTA。'],['关系范式',f'{len(relations)} 条：新增排程、传授、预防、保管、未来代理和通行关系。'],['本轮阅读顺序','先查看 07_本轮变更，再用 01_原创_世界筛选题材、初始状态和长期钩子；绿色为 v0.5，黄色为待生成视觉。'],['版权边界','REF-ONLY 仅限内部抽象研究；不得下载、嵌入、训练、商用或借可识别表达生成原创资产。']]
for r,(k,v) in enumerate(summary,3):
    index.cell(r,1,k).font=Font(bold=True,color='12324A'); index.cell(r,1).fill=accent_fill; index.cell(r,2,v).alignment=Alignment(wrap_text=True,vertical='top'); index.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    for c in range(1,5): index.cell(r,c).border=Border(bottom=thin)
    index.row_dimensions[r].height=44
index['A13']='使用与筛选规则'; index['A13'].font=Font(bold=True,color='FFFFFF'); index['A13'].fill=subheader_fill; index.merge_cells('A13:D13')
notes=[['绿色行','v0.5 新增原创、研究和队列条目；蓝色行保留既有版本。'],['黄色行','PENDING-QUOTA：原创生成 brief 已存在，不代表图像、封面、角色图、地图图或视频已生成。'],['资产标签','ORIG-TEXT、ORIG-AI-VIS、ORIG-DIAGRAM 才进入原创审阅；REF-ONLY 永远不等于可商用。'],['Simulation 用法','将初始状态、资源、地点、人物时钟、关系、事件与状态转场一起运行；不要只拿简介做单线任务。']]
for r,(k,v) in enumerate(notes,14):
    index.cell(r,1,k).font=Font(bold=True); index.cell(r,2,v).alignment=Alignment(wrap_text=True); index.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); index.row_dimensions[r].height=36
for col,w in {'A':22,'B':54,'C':18,'D':18}.items(): index.column_dimensions[col].width=w
index.freeze_panes='A3'
world_new={f'W-{i:02d}' for i in range(39,49)}; char_new={f'C-{i:03d}' for i in range(154,194)}; ip_new={f'IP-{i:02d}' for i in range(38,41)}; vis_new={f'VQ-{i:02d}' for i in range(39,49)}; pending={f'VQ-{i:02d}' for i in range(17,49)}; rel_new={f'RLP-{i:02d}' for i in range(23,29)}
add_sheet('01_原创_世界',['世界 ID','名称','资产属性','题材 / 基调','Simulation 核心','初始状态','核心地点','阵营','主冲突','长期钩子','建议优先级','关联视觉'],worlds,[11,22,16,28,38,39,36,30,35,34,14,15],'OriginalWorlds',new_ids=world_new)
add_sheet('02_原创_角色',['角色 ID','世界 ID','名称','公开身份','公开目标','秘密','资源','底线','关键关系','个人时钟'],characters,[10,10,14,22,25,36,23,27,34,28],'OriginalCharacters',new_ids=char_new)
ipws=add_sheet('03_IP_参考库',['IP ID','IP / 世界','媒介','类型','Simulation 适配度','可借鉴结构','高张力组合','地点 / 阵营 / 事件','官方来源链接','使用标签'],ips,[10,28,18,22,17,38,32,36,50,15],'IPReferences',link_cols={9},new_ids=ip_new)
for r in range(2,len(ips)+2): ipws.cell(r,10).fill=warn_fill; ipws.cell(r,10).font=Font(color='9C1C1C',bold=True)
add_sheet('04_视觉资产',['资产 ID','属性','名称','关联世界','文件 / 来源','规格','视觉关键词','权属标签','使用备注'],visuals,[13,15,28,16,54,20,32,20,45],'VisualAssets',new_ids=vis_new,pending_ids=pending)
add_sheet('05_授权矩阵',['标签','定义','商业使用指引','留档要求'],licenses,[20,48,56,38],'LicenseMatrix')
add_sheet('06_关系范式',['关系 ID','关系范式','适用组合','持续 Simulation 机制','典型演化','设计防呆'],relations,[12,24,36,44,32,42],'RelationshipPatterns',new_ids=rel_new)
add_sheet('07_本轮变更',['版本','日期','类别','变更范围','内容摘要'],change_log,[12,18,16,36,70],'ChangeLog')
for sheet in wb.worksheets: sheet.sheet_properties.pageSetUpPr.fitToPage=True; sheet.page_setup.fitToWidth=1; sheet.page_setup.fitToHeight=0
xlsx=OUT / 'World_Character_素材库_v0.5.xlsx'; wb.save(xlsx)
print(json.dumps(payload['metadata']['counts'],ensure_ascii=False)); print(xlsx); print(OUT / 'world_character_library_data_v0.5.json')
